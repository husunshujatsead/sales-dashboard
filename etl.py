"""
etl.py — One-time conversion of raw Excel sources into tidy Parquet fact tables.

Why this exists
---------------
Reading a 127 MB .xlsm with openpyxl takes minutes. Reading the equivalent
Parquet takes milliseconds. The dashboard must never touch Excel at runtime,
so this script does all the expensive, messy work exactly once and writes
narrow, strongly-typed, categorical-encoded Parquet files.

Run:  python etl.py --raw <folder-with-xlsx> --out data/
"""

from __future__ import annotations

import argparse
import re
import sys
import warnings
from pathlib import Path

import pandas as pd

warnings.filterwarnings("ignore")

# --------------------------------------------------------------------------
# Filename -> loader routing. Matching is on lowercase substrings so users can
# rename files loosely (e.g. "energizer amazon FY26.xlsx") and still be picked up.
# --------------------------------------------------------------------------
PATTERNS = {
    "amazon": ["amazon"],
    "careem": ["careem"],
    "talabat": ["talabat"],
    "sadafco": ["online_shopping", "online shopping"],
    "friesland": ["mss_trend", "mss trend", "base-sku", "base_sku"],
}


def log(msg: str) -> None:
    print(f"  {msg}", flush=True)


def find(raw: Path, key: str) -> Path | None:
    """Locate the first file in `raw` whose name matches any pattern for `key`."""
    for f in sorted(raw.glob("*.xls*")):
        low = f.name.lower()
        if any(p in low for p in PATTERNS[key]):
            return f
    return None


def month_key(ts: pd.Timestamp) -> str:
    return ts.strftime("%Y-%m")



def _as_date(v):
    """Coerce a header cell to a Timestamp, or None.

    pandas hands back datetime.datetime (not pd.Timestamp) from header=None
    reads, and some sheets carry the month as a string. isinstance checks
    against pd.Timestamp alone silently miss both, which yields an empty
    fact table -- so normalise everything through one funnel.
    """
    import datetime as _dt
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (pd.Timestamp, _dt.datetime, _dt.date)):
        return pd.Timestamp(v).normalize().replace(day=1)
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        try:
            return pd.Timestamp(pd.to_datetime(s)).normalize().replace(day=1)
        except Exception:
            return None
    return None


# ==========================================================================
# SADAFCO  —  Online_Shopping_24-26.xlsx
# Already a flat fact table: Year, Month, DepotName, Channel, Customer, SKU...
# The only real work is typing it correctly and building a date column.
# ==========================================================================
def etl_sadafco(path: Path) -> pd.DataFrame:
    log(f"reading {path.name} ...")
    df = pd.read_excel(path, sheet_name=0, engine="openpyxl")
    df.columns = [str(c).strip() for c in df.columns]

    ren = {
        "Channel Level2": "Channel",
        "Categroy": "Category",  # typo exists in the source file
        "Gross Sales Amount": "Sales",
        "Sales Qty": "Qty",
        "DepotName": "Depot",
        "CustomerName": "Customer",
        "CustomerCode": "CustomerCode",
        "ItemSubGroup": "SubGroup",
        "ItemSubGroupDesc": "SubGroupDesc",
        "Type": "Type",
        "SKU": "SKU",
    }
    df = df.rename(columns={k: v for k, v in ren.items() if k in df.columns})

    df["Year"] = pd.to_numeric(df["Year"], errors="coerce").astype("Int64")
    df["Month"] = pd.to_numeric(df["Month"], errors="coerce").astype("Int64")
    df = df.dropna(subset=["Year", "Month"])
    df["Date"] = pd.to_datetime(
        dict(year=df["Year"].astype(int), month=df["Month"].astype(int), day=1)
    )
    df["Sales"] = pd.to_numeric(df["Sales"], errors="coerce").fillna(0.0)
    df["Qty"] = pd.to_numeric(df["Qty"], errors="coerce").fillna(0.0)

    for c in ["Depot", "Channel", "Customer", "Type", "Category", "SubGroup", "SubGroupDesc", "SKU"]:
        if c in df.columns:
            df[c] = df[c].astype(str).str.strip().astype("category")

    keep = [c for c in ["Date", "Year", "Month", "Depot", "Channel", "CustomerCode", "Customer",
                        "Type", "Category", "SubGroup", "SubGroupDesc", "SKU", "Sales", "Qty"]
            if c in df.columns]
    df = df[keep]
    df["MonthKey"] = df["Date"].dt.strftime("%Y-%m").astype("category")
    log(f"sadafco -> {len(df):,} rows, {df['Date'].min():%Y-%m} to {df['Date'].max():%Y-%m}")
    return df


# ==========================================================================
# FRIESLAND  —  MSS_Trend_By_cust_by_Base-SKU_April.xlsm (127 MB)
# Layout: rows 1-4 are pivot slicers, row 7 = month header (merged, forward-fill),
# row 8 = measure header (Sales/Volume/Cases), row 9+ = data.
# Cols A-E are the dimensions: CUSTOMER, Country, Brand, SKU_DESC, Category.
#
# read_only=True streaming is essential -- a normal load_workbook on this file
# consumes several GB of RAM.
# ==========================================================================
def etl_friesland(path: Path) -> pd.DataFrame:
    from openpyxl import load_workbook

    log(f"streaming {path.name} ({path.stat().st_size/1e6:.0f} MB) — this takes a few minutes ...")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]

    rows_iter = ws.iter_rows(values_only=True)
    buf = [next(rows_iter) for _ in range(8)]  # rows 1..8

    month_row = list(buf[6])   # row 7 -> "2025-01" repeated across 3 cols
    meas_row = list(buf[7])    # row 8 -> CUSTOMER..Category, then Sales/Volume/Cases

    # Forward-fill the merged month header across its measure columns
    filled, last = [], None
    for v in month_row:
        s = str(v).strip() if v is not None else ""
        if re.fullmatch(r"\d{4}-\d{2}", s):
            last = s
        filled.append(last)
    month_row = filled

    dim_cols, val_cols = [], []
    for i, m in enumerate(meas_row):
        name = str(m).strip() if m is not None else ""
        if i < 5:
            dim_cols.append((i, name))
        elif name in ("Sales", "Volume", "Cases") and month_row[i]:
            val_cols.append((i, month_row[i], name))

    log(f"detected {len(val_cols)//3} month x 3 measures; dims={[d[1] for d in dim_cols]}")

    dim_idx = [i for i, _ in dim_cols]
    dim_names = [n for _, n in dim_cols]
    val_idx = [i for i, _, _ in val_cols]
    val_meta = [(m, k) for _, m, k in val_cols]

    records, n = [], 0
    for row in rows_iter:
        n += 1
        if row is None:
            continue
        key = row[0]
        if key is None or str(key).strip() in ("", "Grand Total", "None"):
            continue
        dims = [row[i] if i < len(row) else None for i in dim_idx]
        vals = [row[i] if i < len(row) else None for i in val_idx]
        if all(v is None for v in vals):
            continue
        records.append(dims + vals)
        if n % 50000 == 0:
            log(f"  ... {n:,} rows scanned, {len(records):,} kept")

    wb.close()
    log(f"scan complete: {n:,} rows -> {len(records):,} data rows")

    wide = pd.DataFrame.from_records(
        records, columns=dim_names + [f"{m}|{k}" for m, k in val_meta]
    )

    id_vars = dim_names
    long = wide.melt(id_vars=id_vars, var_name="mk", value_name="val")
    long = long.dropna(subset=["val"])
    long["val"] = pd.to_numeric(long["val"], errors="coerce")
    long = long.dropna(subset=["val"])
    long[["MonthKey", "Measure"]] = long["mk"].str.split("|", expand=True)
    long = long.drop(columns=["mk"])

    df = long.pivot_table(
        index=id_vars + ["MonthKey"], columns="Measure", values="val", aggfunc="sum"
    ).reset_index()
    df.columns.name = None

    for c in ("Sales", "Volume", "Cases"):
        if c not in df.columns:
            df[c] = 0.0
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)

    df["Date"] = pd.to_datetime(df["MonthKey"] + "-01")
    df["Year"] = df["Date"].dt.year.astype("int16")
    df["Month"] = df["Date"].dt.month.astype("int8")

    ren = {"CUSTOMER": "Customer", "SKU_DESC": "SKU"}
    df = df.rename(columns={k: v for k, v in ren.items() if k in df.columns})
    for c in ["Customer", "Country", "Brand", "SKU", "Category"]:
        if c in df.columns:
            df[c] = df[c].astype(str).str.strip().astype("category")
    df["MonthKey"] = df["MonthKey"].astype("category")

    keep = [c for c in ["Date", "Year", "Month", "MonthKey", "Customer", "Country",
                        "Brand", "Category", "SKU", "Sales", "Volume", "Cases"] if c in df.columns]
    df = df[keep]
    df = df[(df[["Sales", "Volume", "Cases"]].abs().sum(axis=1) > 0)]
    log(f"friesland -> {len(df):,} tidy rows, {df['Date'].min():%Y-%m} to {df['Date'].max():%Y-%m}")
    return df


# ==========================================================================
# ENERGIZER — three separate retailer files, each a different shape.
# Unified into one fact table with a `Retailer` column.
# ==========================================================================
def _energizer_amazon(path: Path) -> pd.DataFrame:
    """
    Sheet 'v2': row2 = month dates (merged over 3 cols), row3 = measure names
    (Shipped Revenue / Shipped COGS / Shipped Units), row4+ = data.
    Cols A-E = ASIN, BARCODE, Brand, Type, TITLE.
    """
    raw = pd.read_excel(path, sheet_name="v2", header=None, engine="openpyxl")
    mrow, hrow = raw.iloc[1], raw.iloc[2]

    months, last = [], None
    for v in mrow:
        d = _as_date(v)
        if d is not None:
            last = d
        months.append(last)

    recs = []
    body = raw.iloc[3:].reset_index(drop=True)
    body = body[body[0].notna()]
    meas_map = {"Shipped Revenue": "Sales", "Shipped Units": "Units", "Shipped COGS": "COGS"}

    dims = body.iloc[:, :5].copy()
    dims.columns = ["ASIN", "Barcode", "Brand", "Type", "Title"]

    for ci in range(5, raw.shape[1]):
        m, h = months[ci], str(hrow[ci]).strip()
        if m is None or h not in meas_map:
            continue
        vals = pd.to_numeric(body.iloc[:, ci], errors="coerce")
        blk = dims.copy()
        blk["Date"] = m
        blk["Measure"] = meas_map[h]
        blk["val"] = vals
        recs.append(blk)

    if not recs:
        return pd.DataFrame()
    long = pd.concat(recs, ignore_index=True).dropna(subset=["val"])
    df = long.pivot_table(
        index=["ASIN", "Barcode", "Brand", "Type", "Title", "Date"],
        columns="Measure", values="val", aggfunc="sum"
    ).reset_index()
    df.columns.name = None
    df["Retailer"] = "Amazon"
    return df


def _energizer_careem(path: Path) -> pd.DataFrame:
    """Sheets 'Sales' (AED) and 'Volume' (units): row2 = month dates, col A/B = Barcode/Title."""
    out = []
    for sheet, meas in (("Sales", "Sales"), ("Volume", "Units")):
        try:
            raw = pd.read_excel(path, sheet_name=sheet, header=None, engine="openpyxl")
        except Exception:
            continue
        hdr = raw.iloc[1]
        body = raw.iloc[2:].reset_index(drop=True)
        body = body[body[0].notna()]
        for ci in range(2, raw.shape[1]):
            v = _as_date(hdr[ci])
            if v is None:
                continue
            vals = pd.to_numeric(body.iloc[:, ci], errors="coerce")
            blk = pd.DataFrame({
                "Barcode": body.iloc[:, 0].values,
                "Title": body.iloc[:, 1].astype(str).values,
                "Date": v,
                "Measure": meas,
                "val": vals.values,
            })
            out.append(blk)
    if not out:
        return pd.DataFrame()
    long = pd.concat(out, ignore_index=True).dropna(subset=["val"])
    df = long.pivot_table(index=["Barcode", "Title", "Date"], columns="Measure",
                          values="val", aggfunc="sum").reset_index()
    df.columns.name = None
    df["Retailer"] = "Careem"
    df["Brand"] = df["Title"].str.split().str[0]
    return df


def _energizer_talabat(path: Path) -> pd.DataFrame:
    """Small sheet: row2 has month dates every 2 cols (value, units)."""
    try:
        raw = pd.read_excel(path, sheet_name="Sales", header=None, engine="openpyxl")
    except Exception:
        return pd.DataFrame()

    hdr = raw.iloc[1]
    months, last = [], None
    for v in hdr:
        d = _as_date(v)
        if d is not None:
            last = d
        months.append(last)

    body = raw.iloc[2:].reset_index(drop=True)
    label_col = None
    for c in range(min(3, raw.shape[1])):
        if body.iloc[:, c].notna().sum() > 0:
            label_col = c
            break
    if label_col is None:
        return pd.DataFrame()
    body = body[body.iloc[:, label_col].notna()]
    if body.empty:
        return pd.DataFrame()

    recs = []
    seen: dict = {}
    for ci in range(label_col + 1, raw.shape[1]):
        m = months[ci]
        if m is None:
            continue
        k = seen.get(m, 0)
        seen[m] = k + 1
        meas = "Sales" if k == 0 else "Units"
        vals = pd.to_numeric(body.iloc[:, ci], errors="coerce")
        recs.append(pd.DataFrame({
            "Title": body.iloc[:, label_col].astype(str).values,
            "Date": m, "Measure": meas, "val": vals.values,
        }))
    if not recs:
        return pd.DataFrame()
    long = pd.concat(recs, ignore_index=True).dropna(subset=["val"])
    df = long.pivot_table(index=["Title", "Date"], columns="Measure",
                          values="val", aggfunc="sum").reset_index()
    df.columns.name = None
    df["Retailer"] = "Talabat"
    df["Brand"] = df["Title"].str.split().str[0]
    return df


def etl_energizer(raw: Path) -> pd.DataFrame:
    parts = []
    for key, fn in (("amazon", _energizer_amazon),
                    ("careem", _energizer_careem),
                    ("talabat", _energizer_talabat)):
        p = find(raw, key)
        if p is None:
            log(f"energizer: no {key} file found, skipping")
            continue
        log(f"reading {p.name} ...")
        try:
            d = fn(p)
            if len(d):
                parts.append(d)
                log(f"  {key} -> {len(d):,} rows")
        except Exception as e:
            log(f"  !! {key} failed: {e}")

    if not parts:
        return pd.DataFrame()

    df = pd.concat(parts, ignore_index=True)
    for c in ("Sales", "Units", "COGS"):
        if c not in df.columns:
            df[c] = 0.0
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)

    df["Date"] = pd.to_datetime(df["Date"])
    df["Year"] = df["Date"].dt.year.astype("int16")
    df["Month"] = df["Date"].dt.month.astype("int8")
    df["MonthKey"] = df["Date"].dt.strftime("%Y-%m")

    if "Title" not in df.columns:
        df["Title"] = "Unknown"
    df["Title"] = df["Title"].astype(str).str.strip()
    for c in ("Brand", "Type", "ASIN"):
        if c not in df.columns:
            df[c] = "Unknown"
        df[c] = df[c].astype(str).replace({"nan": "Unknown", "None": "Unknown"}).str.strip()

    # Margin only exists where COGS was reported (Amazon)
    df["GrossProfit"] = (df["Sales"] - df["COGS"]).where(df["COGS"] > 0, 0.0)

    for c in ("Retailer", "Brand", "Type", "Title", "MonthKey", "ASIN"):
        df[c] = df[c].astype("category")

    keep = ["Date", "Year", "Month", "MonthKey", "Retailer", "Brand", "Type",
            "ASIN", "Title", "Sales", "Units", "COGS", "GrossProfit"]
    df = df[[c for c in keep if c in df.columns]]
    df = df[df[["Sales", "Units"]].abs().sum(axis=1) > 0]
    log(f"energizer -> {len(df):,} tidy rows, {df['Date'].min():%Y-%m} to {df['Date'].max():%Y-%m}")
    return df


# ==========================================================================
def main() -> int:
    ap = argparse.ArgumentParser(description="Build Parquet fact tables from raw Excel.")
    ap.add_argument("--raw", default="raw", help="folder containing the source .xlsx/.xlsm files")
    ap.add_argument("--out", default="data", help="output folder for .parquet files")
    ap.add_argument("--only", default=None, help="run one only: sadafco|friesland|energizer")
    args = ap.parse_args()

    raw, out = Path(args.raw), Path(args.out)
    if not raw.exists():
        print(f"ERROR: raw folder '{raw}' does not exist.")
        return 1
    out.mkdir(parents=True, exist_ok=True)

    jobs = {
        "sadafco":   lambda: etl_sadafco(find(raw, "sadafco")),
        "friesland": lambda: etl_friesland(find(raw, "friesland")),
        "energizer": lambda: etl_energizer(raw),
    }
    if args.only:
        jobs = {args.only: jobs[args.only]}

    for name, fn in jobs.items():
        print(f"\n[{name}]")
        if name != "energizer" and find(raw, name) is None:
            log("source file not found — skipping")
            continue
        try:
            df = fn()
        except Exception as e:
            log(f"FAILED: {e}")
            continue
        if df is None or df.empty:
            log("no data produced — skipping")
            continue
        dest = out / f"{name}.parquet"
        df.to_parquet(dest, index=False, compression="zstd")
        log(f"wrote {dest}  ({dest.stat().st_size/1e6:.1f} MB)")

    print("\nDone.")
    return 0


if __name__ == "__main__":
    sys.exit(main())